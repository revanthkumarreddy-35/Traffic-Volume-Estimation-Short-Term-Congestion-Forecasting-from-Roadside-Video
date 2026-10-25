"""Apply the supplied Array article's layout to the current manuscript.

Preserves body XML, including equations and images. Uses Times New Roman as
the locally available serif substitute for the reference's Charis SIL.
"""
from copy import deepcopy
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Documents/Paper_Sources/Research paper_updated_authors.docx'
TARGET = ROOT / 'Documents/Research paper.docx'
WIDTH = 520.1
COLUMN = (WIDTH - 18) / 2


def element(name, **attrs):
    node = OxmlElement('w:' + name)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def font(run, size=8, bold=None, italic=None):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    run.font.color.rgb = __import__('docx').shared.RGBColor(0, 0, 0)


def para(p, size=8, before=0, after=3, indent=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing = 1.05
    fmt.first_line_indent = Pt(indent)
    fmt.left_indent = fmt.right_indent = Pt(0)
    fmt.alignment = align
    fmt.keep_with_next = False
    fmt.page_break_before = False
    fmt.widow_control = True
    for run in p.runs: font(run, size)


def border(p, side='bottom', color='888888'):
    borders = p._p.get_or_add_pPr().find(qn('w:pBdr'))
    if borders is None:
        borders = element('pBdr')
        p._p.get_or_add_pPr().append(borders)
    borders.append(element(side, val='single', sz=4, space=5, color=color))


def section(base, columns):
    node = deepcopy(base)
    for name in ('cols', 'type', 'titlePg'):
        for child in list(node.findall(qn('w:' + name))): node.remove(child)
    node.append(element('type', val='continuous'))
    node.append(element('cols', num=columns, space=360, equalWidth=1))
    return node


def section_break_before(anchor, base, columns):
    p = element('p')
    props = element('pPr')
    props.append(element('spacing', before=0, after=0, line=20, lineRule='exact'))
    props.append(section(base, columns))
    p.append(props)
    anchor.addprevious(p)


def style_table(table, width):
    table.autofit = False
    n = len(table.columns)
    widths = [width/n] * n
    if n == 2: widths = [width*.25, width*.75]
    elif n == 3: widths = [width*.46, width*.27, width*.27]
    elif n == 5: widths = [width*.36] + [width*.16]*4
    elif n == 7: widths = [width*x for x in (.08,.10,.10,.26,.20,.14,.12)]
    for col, size in zip(table.columns, widths): col.width = Pt(size)
    props = table._tbl.tblPr
    for old in list(props.findall(qn('w:tblBorders'))): props.remove(old)
    borders = element('tblBorders')
    for side in ('top', 'bottom'):
        borders.append(element(side, val='single', sz=5, color='555555'))
    for side in ('left', 'right', 'insideH', 'insideV'):
        borders.append(element(side, val='nil'))
    props.append(borders)
    margins = element('tblCellMar')
    for side in ('top', 'bottom', 'left', 'right'):
        margins.append(element(side, w=45, type='dxa'))
    props.append(margins)
    for i, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(element('cantSplit'))
        if i == 0: trpr.append(element('tblHeader'))
        for cell, size in zip(row.cells, widths):
            cell.width = Pt(size)
            for p in cell.paragraphs:
                para(p, size=7.2, indent=0, after=1, align=WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.keep_with_next = i < len(table.rows) - 1
                for run in p.runs: font(run, 7.2, bold=(i == 0))
            if i == 0:
                tc_borders = element('tcBorders')
                tc_borders.append(element('bottom', val='single', sz=4, color='888888'))
                cell._tc.get_or_add_tcPr().append(tc_borders)


def main():
    doc = Document(SOURCE)
    source_math = len(doc._element.xpath('.//m:oMath'))
    source_images = len(doc.inline_shapes)
    original_tables = list(doc.tables)
    original_text = [p.text for p in doc.paragraphs if p.text]
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(595.276), Pt(793.701)
    sec.left_margin = sec.right_margin = Pt(37.6)
    sec.top_margin, sec.bottom_margin = Pt(38), Pt(34)
    sec.header_distance, sec.footer_distance = Pt(16), Pt(16)
    sec.different_first_page_header_footer = True
    for name in ('Normal', 'Body Text'):
        doc.styles[name].font.name = 'Times New Roman'
        doc.styles[name].font.size = Pt(8)
    header = sec.header.paragraphs[0]
    header.text = 'P. Charishma et al.\tRoadside traffic counting and congestion classification'
    para(header, size=6.4, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    header.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH), __import__('docx').enum.text.WD_TAB_ALIGNMENT.RIGHT)
    for r in header.runs: r.italic = True
    for footer in (sec.footer, sec.first_page_footer):
        p = footer.paragraphs[0]
        p.clear()
        para(p, 7, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        field = element('fldSimple', instr='PAGE')
        p._p.append(field)
    sec.first_page_header.paragraphs[0].text = ''
    base = deepcopy(sec._sectPr)
    # Initial full-width title/authors/affiliation and abstract block.
    paragraphs = list(doc.paragraphs)
    title = paragraphs[0]
    keywords = next(p for p in paragraphs if p.text.startswith('Keywords:'))
    abstract_heading = next(p for p in paragraphs if p.text == 'Abstract')
    abstract = paragraphs[paragraphs.index(abstract_heading) + 1]
    intro = next(p for p in paragraphs if p.text.startswith('1. Introduction'))
    for p in paragraphs:
        para(p)
        text = p.text.strip()
        if re.match(r'^\d+\.\s', text) or text in ('Declarations', 'References') or text.startswith('Appendix A.'):
            para(p, 8, before=10, after=6, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.keep_with_next = True
            for run in p.runs: run.bold = True
        elif re.match(r'^\d+\.\d+', text):
            para(p, 8, before=8, after=5, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.keep_with_next = True
            for run in p.runs: run.italic = True; run.bold = False
        elif text.startswith(('Table ', 'Figure ', 'Algorithm ')):
            para(p, 7.2, before=4, after=4, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.keep_with_next = text.startswith(('Table ', 'Algorithm '))
        elif re.match(r'^\[\d+\]', text):
            para(p, 6.4, after=2, indent=-9)
            p.paragraph_format.left_indent = Pt(9)
        elif p._p.xpath('.//m:oMath'):
            para(p, 8, before=4, after=4, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        if p._p.xpath('.//w:drawing'):
            para(p, 8, before=5, after=3, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.keep_with_next = True
        if text.startswith('Input: CCTV clips'):
            para(p, 7.2, before=3, after=5, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    # A manuscript masthead uses the reference's restrained grey strip.
    banner = title.insert_paragraph_before('RESEARCH ARTICLE')
    para(banner, 9, before=12, after=14, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    banner._p.get_or_add_pPr().append(element('shd', fill='E7E7E7', val='clear'))
    border(banner)
    para(title, 13.4, before=8, after=8, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
    for run in title.runs: font(run, 13.4, bold=False)
    for p in paragraphs:
        if p.text.startswith('Authors:'):
            p.text = p.text.removeprefix('Authors: ').title()
            para(p, 10.6, after=4, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif p.text.startswith('Guide:'):
            p.text = 'Guide: Dr. T. Praveen'
            para(p, 9, after=4, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif p.text.startswith('Affiliation:'):
            p.text = p.text.removeprefix('Affiliation: ')
            para(p, 7.2, after=8, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT)
            for run in p.runs: run.italic = True
    # The abstract occupies the full article width; omit the keywords sidebar.
    front = doc.add_table(rows=1, cols=1)
    front.autofit = False
    front.columns[0].width = Pt(WIDTH)
    abstract_cell = front.cell(0,0)
    abstract_cell.width = Pt(WIDTH)
    abstract_cell.paragraphs[0].text = 'A B S T R A C T'
    para(abstract_cell.paragraphs[0],8,after=8,indent=0,align=WD_ALIGN_PARAGRAPH.LEFT)
    abstract_cell.paragraphs[0].paragraph_format.keep_with_next = True
    border(abstract_cell.paragraphs[0])
    p = abstract_cell.add_paragraph(abstract.text)
    para(p, 8, after=6, indent=0)
    keywords._p.addprevious(front._tbl)
    for p in (keywords, abstract_heading, abstract): p._p.getparent().remove(p._p)
    # Preserve the paper's existing images and equations, adapting widths.
    for shape in doc.inline_shapes:
        ratio = shape.height / shape.width
        image_p = next(node for node in shape._inline.iterancestors() if node.tag == qn('w:p'))
        caption = image_p.getnext()
        caption_text = ''.join(caption.itertext()) if caption is not None else ''
        size = COLUMN if 'Detection-confidence bands' in caption_text else WIDTH
        shape.width = Pt(size)
        shape.height = Pt(size * ratio)
    # Body uses two columns, with continuous full-width sections around large exhibits.
    section_break_before(intro._p, base, 1)
    for table_index, table in enumerate(original_tables):
        wide = len(table.columns) >= 5 and table_index != 1
        style_table(table, WIDTH if wide else COLUMN)
        if wide:
            prev = table._tbl.getprevious()
            # Keep its caption in the full-width section.
            anchor = prev if prev is not None and prev.tag == qn('w:p') else table._tbl
            section_break_before(anchor, base, 2)
            next_node = table._tbl.getnext()
            if next_node is not None: section_break_before(next_node, base, 1)
    for p in list(doc.paragraphs):
        if not p._p.xpath('.//w:drawing'): continue
        caption = p._p.getnext()
        if caption is not None and 'Detection-confidence bands' in ''.join(caption.itertext()):
            continue
        section_break_before(p._p, base, 2)
        caption = p._p.getnext()
        after = caption.getnext() if caption is not None else None
        if after is not None: section_break_before(after, base, 1)
    final = doc._element.body.sectPr
    doc._element.body.replace(final, section(base, 2))
    # Keep page numbering on the first page too; use distinct blank first header.
    first_break = intro._p.getprevious().find(qn('w:pPr')).find(qn('w:sectPr'))
    first_break.append(element('titlePg'))
    doc.save(TARGET)
    check = Document(TARGET)
    assert len(check.inline_shapes) == source_images
    assert len(check._element.xpath('.//m:oMath')) == source_math
    assert len(check.tables) == len(original_tables) + 1
    texts = [p.text for p in check.paragraphs]
    exceptions = ('Authors:', 'Guide:', 'Affiliation:', 'Keywords:')
    for text in original_text:
        if text == 'Abstract' or text == abstract.text or text.startswith(exceptions): continue
        assert text in texts, 'Content lost: ' + text[:80]
    print('Created:', TARGET)
    print('Preserved', source_images, 'figures,', source_math, 'equations and', len(original_tables), 'research tables.')


if __name__ == '__main__': main()

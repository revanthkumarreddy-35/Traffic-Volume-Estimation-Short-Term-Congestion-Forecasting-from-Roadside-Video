"""Add verified reference-source hyperlinks to the styled journal manuscript."""
import csv
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / 'Documents/Research paper.docx'
# Paper titles/authors checked against primary sources on 23 October 2026.
# DOI endpoints for 7 and 22 restrict automated retrieval; their identifiers
# were corroborated against the chapter/research records and publisher citations.
LINKS = {
    1: 'https://arxiv.org/abs/1506.02640',
    2: 'https://docs.ultralytics.com/models/yolo11/',
    3: 'https://arxiv.org/abs/1511.04136',
    4: 'https://arxiv.org/abs/1602.00763',
    5: 'https://arxiv.org/abs/1703.07402',
    6: 'https://arxiv.org/abs/2110.06864',
    7: 'https://doi.org/10.1016/B978-0-12-336156-1.50061-6',
    8: 'https://arxiv.org/abs/1603.02754',
    9: 'https://doi.org/10.48084/etasr.18369',
    10: 'https://arxiv.org/abs/2604.04080',
    11: 'https://arxiv.org/abs/1804.02767',
    12: 'https://arxiv.org/abs/1506.01497',
    13: 'https://arxiv.org/abs/1512.02325',
    14: 'https://arxiv.org/abs/2207.02696',
    15: 'https://arxiv.org/abs/2410.17725',
    16: 'https://arxiv.org/abs/2206.14651',
    17: 'https://arxiv.org/abs/2203.14360',
    18: 'https://arxiv.org/abs/2202.13514',
    19: 'https://doi.org/10.1016/S0734-189X(87)80186-X',
    20: 'https://arxiv.org/abs/1405.0312',
    21: 'https://doi.org/10.1109/TITS.2014.2345663',
    22: 'https://doi.org/10.1061/(ASCE)0733-947X(2003)129:6(664)',
    23: 'https://doi.org/10.1162/neco.1997.9.8.1735',
    24: 'https://www.techscience.com/cmc/v86n1/64404/html',
    25: 'https://ejournal.unibabwi.ac.id/index.php/Zetroem/article/view/8313',
    26: 'https://proceedings.neurips.cc/paper_files/paper/2017/hash/6449f44a102fde848669bdd9eb6b76fa-Abstract.html',
}


def main():
    doc = Document(PAPER)
    image_count = len(doc.inline_shapes)
    equation_count = len(doc._element.xpath('.//m:oMath'))
    found = set()
    records = []
    for paragraph in doc.paragraphs:
        match = re.match(r'^\[(\d+)\]', paragraph.text)
        if not match:
            continue
        number = int(match.group(1))
        url = LINKS[number]
        # Reruns replace only links created by this script.
        for link in list(paragraph._p.findall(qn('w:hyperlink'))):
            if ''.join(link.itertext()).strip().endswith('[Source]'):
                paragraph._p.remove(link)
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
        link.set(qn('w:history'), '1')
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        for tag, attrs in [('rFonts', {'ascii': 'Times New Roman', 'hAnsi': 'Times New Roman'}),
                           ('sz', {'val': '13'}), ('color', {'val': '007AAB'}),
                           ('u', {'val': 'single'})]:
            node = OxmlElement('w:' + tag)
            for key, value in attrs.items(): node.set(qn('w:' + key), value)
            props.append(node)
        run.append(props)
        label = OxmlElement('w:t')
        label.set(qn('xml:space'), 'preserve')
        label.text = ' [Source]'
        run.append(label)
        link.append(run)
        paragraph._p.append(link)
        found.add(number)
        records.append({'reference': number, 'citation': paragraph.text.removesuffix(' [Source]'), 'url': url})
    assert found == set(LINKS), 'Missing bibliography entries'
    pending = PAPER.with_name('Research paper_links_pending.docx')
    doc.save(pending)
    checked = Document(pending)
    for paragraph in checked.paragraphs:
        match = re.match(r'^\[(\d+)\]', paragraph.text)
        if match:
            expected = LINKS[int(match.group(1))]
            targets = [paragraph.part.rels[link.get(qn('r:id'))].target_ref
                       for link in paragraph._p.findall(qn('w:hyperlink'))]
            assert expected in targets
    assert len(checked.inline_shapes) == image_count
    assert len(checked._element.xpath('.//m:oMath')) == equation_count
    pending.replace(PAPER)
    with (ROOT / 'CP-1/Executed_Results_Analysis/reference_links.csv').open('w', encoding='utf-8', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=['reference', 'citation', 'url'])
        writer.writeheader()
        writer.writerows(records)
    print('Added 26 clickable reference links:', PAPER)


if __name__ == '__main__':
    main()

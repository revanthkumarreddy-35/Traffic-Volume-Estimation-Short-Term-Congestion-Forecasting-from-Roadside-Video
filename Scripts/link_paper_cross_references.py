"""Link in-text citations and figure/table mentions to internal DOCX bookmarks."""
from collections import Counter
from copy import deepcopy
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / 'Documents/Research paper.docx'
PREFIX = 'cp1_'


def bookmark(paragraph_element, name, number):
    start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    start.set(qn('w:id'), str(number))
    start.set(qn('w:name'), name)
    end.set(qn('w:id'), str(number))
    index = 1 if paragraph_element.find(qn('w:pPr')) is not None else 0
    paragraph_element.insert(index, start)
    paragraph_element.append(end)


def copy_text_run(source_run, text):
    result = OxmlElement('w:r')
    props = source_run.find(qn('w:rPr'))
    if props is not None:
        result.append(deepcopy(props))
    Run(result, None).text = text
    return result


def link_span(paragraph, start, end, destination):
    """Split only affected runs, retaining their font/style and surrounding XML."""
    runs = []
    offset = 0
    for run in paragraph.runs:
        size = len(run.text)
        if offset < end and offset + size > start:
            runs.append((run, offset, offset + size))
        offset += size
    if not runs:
        raise ValueError('Unable to locate cross-reference text')
    first, first_offset, _ = runs[0]
    last, last_offset, _ = runs[-1]
    insert = []
    prefix = first.text[:start-first_offset]
    if prefix:
        insert.append(copy_text_run(first._r, prefix))
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), destination)
    hyperlink.set(qn('w:history'), '1')
    for run, low, high in runs:
        text = run.text[max(0, start-low):min(high-low, end-low)]
        linked_run = copy_text_run(run._r, text)
        props = linked_run.find(qn('w:rPr'))
        if props is None:
            props = OxmlElement('w:rPr')
            linked_run.insert(0, props)
        for node in list(props.findall(qn('w:color'))):
            props.remove(node)
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '007AAB')
        props.append(color)
        hyperlink.append(linked_run)
    insert.append(hyperlink)
    suffix = last.text[end-last_offset:]
    if suffix:
        insert.append(copy_text_run(last._r, suffix))
    for node in insert:
        first._r.addprevious(node)
    for run, _, _ in runs:
        run._r.getparent().remove(run._r)


def main():
    doc = Document(PAPER)
    original_text = [p.text for p in doc.paragraphs]
    equations = len(doc._element.xpath('.//m:oMath'))
    images = len(doc.inline_shapes)
    external_links = len(doc._element.xpath('.//w:hyperlink[@r:id]'))
    # Remove only generated bookmarks/links so reruns do not add duplicates.
    for link in list(doc._element.xpath('.//w:hyperlink[@w:anchor]')):
        if link.get(qn('w:anchor'), '').startswith(PREFIX):
            for child in list(link):
                link.addprevious(child)
            link.getparent().remove(link)
    owned_ids = {b.get(qn('w:id')) for b in doc._element.xpath('.//w:bookmarkStart')
                 if b.get(qn('w:name'), '').startswith(PREFIX)}
    for tag in ('bookmarkStart', 'bookmarkEnd'):
        for node in list(doc._element.xpath('.//w:' + tag)):
            if node.get(qn('w:id')) in owned_ids:
                node.getparent().remove(node)
    next_id = max([int(b.get(qn('w:id'))) for b in doc._element.xpath('.//w:bookmarkStart')] + [0]) + 1
    destinations = {}
    for paragraph in doc.paragraphs:
        ref = re.match(r'^\[(\d+)\]', paragraph.text)
        caption = re.match(r'^(Figure|Table)\s+(A?\d+)\.', paragraph.text)
        if ref:
            key = 'ref_' + ref.group(1)
        elif caption:
            key = caption.group(1).lower() + '_' + caption.group(2)
        else:
            continue
        target = paragraph._p
        if caption and caption.group(1) == 'Figure':
            previous = target.getprevious()
            if previous is not None and previous.xpath('.//w:drawing'):
                target = previous  # Include the picture in the jump viewport.
        name = PREFIX + key
        if name in destinations:
            raise ValueError('Duplicate destination: ' + name)
        bookmark(target, name, next_id)
        next_id += 1
        destinations[name] = paragraph.text
    # New evaluated experiments may add exhibits; require continuous numbering.
    assert {PREFIX+'ref_'+str(i) for i in range(1,27)} <= set(destinations)
    assert PREFIX+'table_A1' in destinations
    for kind in ('figure','table'):
        numbers = sorted(int(name.rsplit('_',1)[1]) for name in destinations
                         if re.fullmatch(PREFIX+kind+r'_\d+',name))
        assert numbers == list(range(1,max(numbers)+1)), 'Non-contiguous '+kind+' numbering'
    counts = Counter()
    records = []
    for paragraph in doc.paragraphs:
        if re.match(r'^(?:\[\d+\]|(?:Figure|Table)\s+A?\d+\.)', paragraph.text):
            continue
        text = ''.join(run.text for run in paragraph.runs)
        spans = []
        for match in re.finditer(r'\[(\d+)\]', text):
            spans.append((match.start(), match.end(), PREFIX+'ref_'+match.group(1)))
        # Link visible endpoints individually in ranges such as Tables 2-3, 5-6.
        for match in re.finditer(r'\b(Figures?|Tables?)\s+(A?\d+(?:\s*[-–,]\s*A?\d+)*)', text):
            kind = match.group(1).rstrip('s').lower()
            numbers = list(re.finditer(r'A?\d+', match.group(2)))
            if len(numbers) == 1:
                spans.append((match.start(), match.end(), PREFIX+kind+'_'+numbers[0].group()))
            else:
                for number in numbers:
                    spans.append((match.start(2)+number.start(), match.start(2)+number.end(), PREFIX+kind+'_'+number.group()))
        # Right-to-left edits leave positions of preceding plain runs unchanged.
        for start, end, target in sorted(spans, reverse=True):
            if target not in destinations:
                raise ValueError('Missing destination: ' + target)
            link_span(paragraph, start, end, target)
            counts[target.removeprefix(PREFIX).split('_')[0]] += 1
            records.append({'text': text[start:end], 'target': target})
    assert [p.text for p in doc.paragraphs] == original_text
    assert len(doc.inline_shapes) == images
    assert len(doc._element.xpath('.//m:oMath')) == equations
    assert len(doc._element.xpath('.//w:hyperlink[@r:id]')) == external_links
    pending = PAPER.with_name('Research paper_crosslinks_pending.docx')
    doc.save(pending)
    check = Document(pending)
    targets = {b.get(qn('w:name')) for b in check._element.xpath('.//w:bookmarkStart')}
    generated = [h for h in check._element.xpath('.//w:hyperlink[@w:anchor]')
                 if h.get(qn('w:anchor')).startswith(PREFIX)]
    assert len(generated) == sum(counts.values())
    assert all(h.get(qn('w:anchor')) in targets for h in generated)
    pending.replace(PAPER)
    report = {'counts': dict(counts), 'destinations': destinations, 'links': records}
    (ROOT / 'CP-1/Executed_Results_Analysis/internal_cross_references.json').write_text(
        json.dumps(report, indent=2), encoding='utf-8')
    print('Added internal links:', dict(counts))
    print('Retained external source links:', external_links)


if __name__ == '__main__':
    main()
